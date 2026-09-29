#!/usr/bin/env python3
"""Build the five-tab monthly meal-plan workbook from a filled-in plan data file.

    python3 build_workbook.py plan_data.py /mnt/user-data/outputs/Meal_Plan_October_2026.xlsx

The data file must define: META, ORDER1, ORDER2, CALENDAR, RECIPES, COVERAGE,
CARRY_ONCE, READ_ME, RULES, SAVEA, AARON.  See assets/plan_data_template.py.

Every total is written as a live formula over cell references, never a computed
constant, so the user's edits recalculate.  After running this, always:
    python scripts/recalc.py <out> 90        # from the xlsx skill, expect total_errors: 0
    python3 audit_workbook.py <out> plan_data.py
"""
import sys, os, importlib.util
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

F = "Arial"
BLUE   = Font(name=F, size=10, color="0000FF")
BODY   = Font(name=F, size=10)
BOLD   = Font(name=F, size=10, bold=True)
BOLD11 = Font(name=F, size=11, bold=True)
TITLE  = Font(name=F, size=16, bold=True, color="1F3864")
SUB    = Font(name=F, size=11, italic=True, color="404040")
HDRF   = Font(name=F, size=10, bold=True, color="FFFFFF")
GREEN  = Font(name=F, size=10, bold=True, color="006100")

FILL_HDR  = PatternFill("solid", fgColor="1F3864")
FILL_SEC  = PatternFill("solid", fgColor="D9E1F2")
FILL_KID  = PatternFill("solid", fgColor="FFF2CC")
FILL_SPEC = PatternFill("solid", fgColor="E4DFEC")
FILL_PREP = PatternFill("solid", fgColor="E2EFDA")
FILL_FREE = PatternFill("solid", fgColor="F2F2F2")
FILL_ORD  = PatternFill("solid", fgColor="DDEBF7")
FILL_WARN = PatternFill("solid", fgColor="FCE4D6")
FILL_TOT  = PatternFill("solid", fgColor="FFF2CC")

_t = Side(style="thin", color="BFBFBF")
BORD = Border(left=_t, right=_t, top=_t, bottom=_t)
TOP  = Alignment(vertical="top", wrap_text=True)
TOPL = Alignment(vertical="top", wrap_text=True, horizontal="left")
CTR  = Alignment(vertical="center", horizontal="center", wrap_text=True)
FILLMAP = {"kid": FILL_KID, "special": FILL_SPEC, "prep": FILL_PREP,
           "free": FILL_FREE, "order": FILL_ORD, "n": None}


def load(path):
    spec = importlib.util.spec_from_file_location("plan_data", path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules["plan_data"] = mod
    spec.loader.exec_module(mod)
    return mod


def head(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = HDRF, FILL_HDR, CTR, BORD
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def band(ws, row, text, ncols, fill=FILL_SEC):
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.alignment = BOLD11, TOPL
    for i in range(1, ncols + 1):
        ws.cell(row=row, column=i).fill = fill
        ws.cell(row=row, column=i).border = BORD


def build(D, out):
    wb = openpyxl.Workbook()
    M = D.META

    # ---------------- Shopping Lists (built first: Overview links to its rows)
    sl = wb.active
    sl.title = "Shopping Lists"
    sl["A1"] = f"SHOPPING LISTS  |  {M['date_range']}  |  Two Instacart orders"
    sl["A1"].font = TITLE
    sl["A2"] = ("Blue cells are yours to edit - change a price or a quantity and every total "
                "recalculates. Prices already include a ~18% Instacart markup, so curbside "
                "pickup runs noticeably cheaper.")
    sl["A2"].font = SUB
    sl.merge_cells("A2:J2"); sl.row_dimensions[2].height = 28

    head(sl, 4, ["Order", "Item", "Store", "Pack / Size", "Priority", "Unit Price", "Qty",
                 "Line Total", "Used for", "Notes"],
         [9, 34, 11, 15, 10, 11, 6, 12, 42, 56])

    r = first = 5
    for label, rows in ((M["order1_label"], D.ORDER1), (M["order2_label"], D.ORDER2)):
        band(sl, r, label, 10); r += 1
        for o, item, store, pack, pri, price, qty, used, note in rows:
            for i, v in enumerate([o, item, store, pack, pri], 1):
                c = sl.cell(row=r, column=i, value=v)
                c.font = BOLD if (i == 5 and pri == "Core") else BODY
            pc = sl.cell(row=r, column=6, value=price); pc.font = BLUE
            pc.number_format = '$#,##0.00'
            sl.cell(row=r, column=7, value=qty).font = BLUE
            tc = sl.cell(row=r, column=8, value=f"=F{r}*G{r}")
            tc.font, tc.number_format = BODY, '$#,##0.00'
            sl.cell(row=r, column=9, value=used).font = BODY
            sl.cell(row=r, column=10, value=note).font = BODY
            for i in range(1, 11):
                cc = sl.cell(row=r, column=i); cc.alignment, cc.border = TOP, BORD
                if pri == "Optional":
                    cc.fill = FILL_FREE
                if note.startswith("!!"):
                    cc.fill = FILL_WARN
            sl.row_dimensions[r].height = 30
            r += 1
        r += 1
    last = r - 2
    PR, TR, OR_ = f"$E${first}:$E${last}", f"$H${first}:$H${last}", f"$A${first}:$A${last}"

    r += 1
    band(sl, r, "BUDGET  -  live totals, all driven by the blue cells above", 10, FILL_TOT)
    r += 1
    s = r
    budget = M.get("budget", 800)
    for lbl, val in [
        ("Core  (this month's meals + the staple snack bench)", f'=SUMIFS({TR},{PR},"Core")'),
        ("Pantry  (multi-month staples, right-sized - skip what you already own)",
         f'=SUMIFS({TR},{PR},"Pantry")'),
        ("COUNTED TOWARD BUDGET  (Core + Pantry)", f"=H{s}+H{s+1}"),
        ("Monthly budget", budget),
        ("CUSHION REMAINING", f"=H{s+3}-H{s+2}"),
        ("", ""),
        ("Optional  (snack variety + upgrades - NOT counted)", f'=SUMIFS({TR},{PR},"Optional")'),
        ("Grand total if you add everything Optional", f"=H{s+2}+H{s+6}"),
        ("", ""),
        ("Order 1 subtotal (counted)",
         f'=SUMIFS({TR},{OR_},"Order 1")-SUMIFS({TR},{OR_},"Order 1",{PR},"Optional")'),
        ("Order 2 subtotal (counted)",
         f'=SUMIFS({TR},{OR_},"Order 2")-SUMIFS({TR},{OR_},"Order 2",{PR},"Optional")'),
        ("", ""),
        ("If the Pantry block is already stocked, real spend drops to", f"=H{s}"),
        ("If you do curbside instead of delivery (strips the ~18% markup)", f"=H{s+2}/1.18"),
    ]:
        if lbl:
            c = sl.cell(row=r, column=2, value=lbl)
            c.font = BOLD if ("BUDGET" in lbl or "CUSHION" in lbl) else BODY
            c.alignment = TOPL
            v = sl.cell(row=r, column=8, value=val)
            v.number_format = '$#,##0.00'
            v.font = BLUE if lbl == "Monthly budget" else BOLD
            v.border = BORD
            if "CUSHION" in lbl or "COUNTED" in lbl:
                for i in (2, 8):
                    sl.cell(row=r, column=i).fill = FILL_TOT
        r += 1
    B = dict(core=s, pantry=s+1, total=s+2, budget=s+3, cushion=s+4,
             opt=s+6, o1=s+9, o2=s+10, stocked=s+12, curb=s+13)

    r += 1
    sl.cell(row=r, column=2, value=(
        "LEGEND:  Blue = edit me.  Core = the month's meals.  Pantry = multi-month staples, "
        "right-sized (skip if stocked).  Optional (grey rows) = variety and upgrades, "
        "not in the budget total.")).font = SUB
    sl.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
    sl.row_dimensions[r].height = 26
    sl.freeze_panes = "A5"

    # ---------------- Overview
    ov = wb.create_sheet("Overview", 0)
    ov.column_dimensions["A"].width = 3
    for col, w in zip("BCDEF", [30, 26, 26, 34, 46]):
        ov.column_dimensions[col].width = w
    ov["B1"] = M["title"]; ov["B1"].font = TITLE
    ov["B2"] = M["subtitle"]; ov["B2"].font = SUB

    row = 4

    def sec(t):
        nonlocal row
        band(ov, row, t, 6); row += 1

    def kv(k, v, bold=False, fill=None):
        nonlocal row
        a = ov.cell(row=row, column=2, value=k)
        a.font, a.alignment = (BOLD if bold else BODY), TOPL
        b = ov.cell(row=row, column=3, value=v)
        b.font, b.alignment = (BOLD if bold else BODY), TOPL
        ov.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        if fill:
            for i in range(2, 7):
                ov.cell(row=row, column=i).fill = fill
        ov.row_dimensions[row].height = 30
        row += 1

    sec("HOW THIS WORKBOOK IS ORGANIZED")
    for k, v in [
        ("Overview", "This page. Budget at a glance, the order schedule, the coverage check, and the read-me notes."),
        ("Calendar", "Every day in the range: dinner, calories, the prepped lunch, snacks, who's home, and which order supplies it."),
        ("Shopping Lists", "Both Instacart orders, split by store. Blue cells are editable; totals are live formulas."),
        ("Recipes", "Every dinner, lunch, special meal and homemade snack."),
        ("Snack Reference", "The low-cal grab bench (olive-free) and Aaron's fresh + homemade list."),
    ]:
        kv(k, v)

    row += 1
    sec("BUDGET AT A GLANCE  (live - pulls straight from Shopping Lists)")
    for k, ref, note in [
        ("Core", B["core"], "Meals + staple snacks"),
        ("Pantry", B["pantry"], "Multi-month staples, right-sized"),
        ("COUNTED TOWARD BUDGET", B["total"], "Core + Pantry"),
        ("Budget", B["budget"], "Standing monthly target"),
        ("CUSHION REMAINING", B["cushion"], "Under budget by this much"),
        ("Optional (not counted)", B["opt"], "Snack variety + upgrades"),
        ("Order 1 subtotal", B["o1"], M["order1_short"]),
        ("Order 2 subtotal", B["o2"], M["order2_short"]),
        ("...if pantry is stocked", B["stocked"], "Deeper cushion"),
        ("...if you do curbside", B["curb"], "Strips the ~18% Instacart markup"),
    ]:
        a = ov.cell(row=row, column=2, value=k)
        a.font, a.alignment = (BOLD if k.isupper() else BODY), TOPL
        b = ov.cell(row=row, column=3, value=f"='Shopping Lists'!$H${ref}")
        b.number_format, b.font, b.border = '$#,##0.00', GREEN, BORD
        b.alignment = Alignment(vertical="top", horizontal="left")
        c = ov.cell(row=row, column=4, value=note); c.font, c.alignment = SUB, TOPL
        ov.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
        if k.isupper():
            for i in range(2, 7):
                ov.cell(row=row, column=i).fill = FILL_TOT
        row += 1

    row += 1
    sec("ORDER SCHEDULE  -  read this before you cook anything")
    for k, v in M["order_schedule"]:
        kv(k, v)

    row += 1
    sec("!! ORDER COVERAGE CHECK  -  nothing is assumed to be in the house before it is ordered")
    kv("The rule", "A meal never uses a fresh ingredient that has to survive longer than it can. "
                   "Where a date sits far from its delivery, the plan uses frozen, hardy or "
                   "shelf-stable ingredients on purpose.", bold=True, fill=FILL_WARN)
    row += 1
    for i, h in enumerate(["Date", "Meal", "Days since delivery",
                           "What could have gone wrong", "Why it holds"], 2):
        c = ov.cell(row=row, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = HDRF, FILL_HDR, CTR, BORD
    ov.row_dimensions[row].height = 28
    row += 1
    for rec in D.COVERAGE:
        for i, v in enumerate(rec, 2):
            c = ov.cell(row=row, column=i, value=v)
            c.font = BOLD if i == 2 else BODY
            c.alignment, c.border = TOP, BORD
        ov.row_dimensions[row].height = 28
        row += 1

    row += 1
    kv("Bought ONCE and carried across both halves", D.CARRY_ONCE, bold=True, fill=FILL_ORD)
    row += 1
    kv("So if Order 2 does not list something you expect", M["order2_gap_note"], fill=FILL_ORD)

    row += 2
    sec("READ ME  -  what's new this month")
    for k, v in D.READ_ME:
        kv(k, v)

    row += 1
    sec("STANDING RULES  -  checked against every line of this plan")
    for k, v in D.RULES:
        kv(k, v)

    # ---------------- Calendar
    cal = wb.create_sheet("Calendar", 1)
    cal["A1"] = f"CALENDAR  |  {M['date_range']}"; cal["A1"].font = TITLE
    cal["A2"] = ("Calories are Aaron's serving unless noted. Savea's column is her scaled portion. "
                 "'Supplied by' tells you which Instacart order paid for that night.")
    cal["A2"].font = SUB
    cal.merge_cells("A2:K2")
    head(cal, 4, ["Date", "Day", "Who's home", "Dinner", "Aaron kcal", "Savea kcal",
                  "Prep-ahead lunch", "Aaron snack ideas", "Savea snack ideas (NO olives)",
                  "Supplied by", "Notes"],
         [10, 6, 17, 40, 8, 8, 30, 30, 32, 11, 54])
    r = 5
    for rec in D.CALENDAR:
        *vals, typ = rec
        for i, v in enumerate(vals, 1):
            c = cal.cell(row=r, column=i, value=v)
            c.font = BOLD if (i in (1, 4) and typ in ("kid", "special", "order", "prep")) else BODY
            c.alignment = CTR if i in (2, 5, 6, 10) else TOP
            c.border = BORD
            if FILLMAP[typ]:
                c.fill = FILLMAP[typ]
        cal.row_dimensions[r].height = 44
        r += 1
    r += 1
    band(cal, r, "COLOR KEY", 11); r += 1
    for lbl, fill in [("Kid days - Daniel here (Wed dinner through Sat breakfast)", FILL_KID),
                      ("Special meal / guests", FILL_SPEC),
                      ("Prep Sunday", FILL_PREP),
                      ("Leftovers / free", FILL_FREE),
                      ("Instacart order day", FILL_ORD)]:
        c = cal.cell(row=r, column=1, value=""); c.fill, c.border = fill, BORD
        cal.cell(row=r, column=2, value=lbl).font = BODY
        cal.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        r += 1
    cal.freeze_panes = "A5"

    # ---------------- Recipes
    rc = wb.create_sheet("Recipes", 3)
    rc["A1"] = f"RECIPES  |  {M['month_label']}"; rc["A1"].font = TITLE
    rc["A2"] = "Calories and protein are Aaron's serving. Savea's portion runs roughly 65-75% of it."
    rc["A2"].font = SUB
    head(rc, 4, ["Recipe", "Type", "Serves", "kcal", "Protein (g)", "Method",
                 "Key ingredients", "Quick steps", "Order"],
         [38, 14, 7, 7, 10, 20, 46, 76, 10])
    r = 5
    for name, typ, serves, kcal, prot, method, ingr, steps, order in D.RECIPES:
        for i, v in enumerate([name, typ, serves or "", kcal, prot, method, ingr, steps, order], 1):
            c = rc.cell(row=r, column=i, value=v)
            c.font = BOLD if i == 1 else BODY
            c.alignment = CTR if i in (3, 4, 5) else TOP
            c.border = BORD
            if typ == "Special":
                c.fill = FILL_SPEC
            elif typ == "Lunch":
                c.fill = FILL_PREP
            elif typ == "Snack":
                c.fill = FILL_FREE
        rc.row_dimensions[r].height = max(56, 13 * (steps.count("\n") + 1) + 14)
        r += 1
    rc.freeze_panes = "A5"

    # ---------------- Snack Reference
    sn = wb.create_sheet("Snack Reference", 4)
    sn["A1"] = f"SNACK REFERENCE  |  {M['month_label']}"; sn["A1"].font = TITLE
    sn["A2"] = ("Two separate systems. Savea's bench is grab-and-go and completely olive-free. "
                "Aaron's list includes olives and the homemade batch items.")
    sn["A2"].font = SUB
    head(sn, 4, ["Item", "Category", "Calories", "Protein (g)", "Priority", "Notes"],
         [36, 16, 11, 12, 11, 62])
    r = 5

    def snack_block(banner, rows, fill, warn_key=None):
        nonlocal r
        band(sn, r, banner, 6, fill); r += 1
        for item, cat, cal_, prot, pri, note in rows:
            if item.startswith("--"):
                sn.cell(row=r, column=1, value=item.strip("- ").strip()).font = BOLD11
                for i in range(1, 7):
                    sn.cell(row=r, column=i).fill = FILL_SEC
                    sn.cell(row=r, column=i).border = BORD
                r += 1
                continue
            hot = warn_key and warn_key in item
            for i, v in enumerate([item, cat, cal_, prot, pri, note], 1):
                c = sn.cell(row=r, column=i, value=v)
                c.font = BOLD if hot else BODY
                c.alignment = CTR if i in (3, 4, 5) else TOP
                c.border = BORD
                if hot:
                    c.fill = FILL_WARN
                elif pri == "Optional":
                    c.fill = FILL_FREE
            r += 1

    snack_block("SAVEA'S GRAB BENCH  -  NO OLIVES, EVER.  Target ~300-500 cal/day across the "
                "whole bench, zero decisions, zero prep.", D.SAVEA, FILL_WARN)
    r += 1
    snack_block("AARON'S LIST  -  olives welcome.  Roughly 1,800-2,000 cal/day overall, so snacks "
                "carry more weight here.", D.AARON, FILL_ORD, warn_key="KALAMATA")
    sn.freeze_panes = "A5"

    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    return out, B


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit("usage: build_workbook.py <plan_data.py> <output.xlsx>")
    data = load(sys.argv[1])
    path, B = build(data, sys.argv[2])
    print(f"saved {path}")
    print(f"budget block rows: {B}")
    print("NEXT: run recalc.py (expect total_errors: 0), then audit_workbook.py")
