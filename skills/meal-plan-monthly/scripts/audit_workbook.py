#!/usr/bin/env python3
"""Audit a finished meal-plan workbook against every hard rule.

    python3 audit_workbook.py Meal_Plan_October_2026.xlsx [plan_data.py]

Naive keyword searching produces false positives: the Overview states the rule
"NO ORZO" and the Snack Reference banner says "NO OLIVES, EVER", so both contain
the forbidden word.  This script scopes each check to the cells where a real
violation would actually live, and reports rule-statement hits separately so
they can be confirmed rather than silently dismissed.

Exit code 0 = all pass, 1 = at least one FAIL.
"""
import sys, re, collections
import openpyxl

OK, BAD = "PASS", "FAIL"
results = []


def rec(name, status, detail=""):
    results.append((name, status, detail))


def txt(c):
    return c.value if isinstance(c.value, str) else ""


def audit(path, data_path=None):
    wb = openpyxl.load_workbook(path)
    names = wb.sheetnames

    # --- 1. five tabs, correctly named -------------------------------------
    want = ["Overview", "Calendar", "Shopping Lists", "Recipes", "Snack Reference"]
    rec("Five tabs, correct names", OK if names == want else BAD, str(names))
    if not all(n in names for n in want):
        return finish()
    cal, sl, rc, sn, ov = (wb["Calendar"], wb["Shopping Lists"], wb["Recipes"],
                           wb["Snack Reference"], wb["Overview"])

    cal_rows = []                     # stop at the first blank date: below it is the colour key
    for r in range(5, cal.max_row + 1):
        if not txt(cal.cell(row=r, column=1)):
            break
        cal_rows.append(r)
    # shopping item rows are exactly those carrying a real priority; this excludes the
    # section bands, the budget block and the legend, which otherwise read as violations
    sl_rows = [r for r in range(5, sl.max_row + 1)
               if txt(sl.cell(row=r, column=5)) in ("Core", "Pantry", "Optional")]

    # --- 2. NO ORZO in any actual food content ------------------------------
    food_hits, stmt_hits = [], []
    for ws, cols in ((cal, [4, 7, 8, 9]), (rc, [1, 7, 8]), (sl, [2, 9])):
        for r in range(1, ws.max_row + 1):
            for c in cols:
                if "orzo" in txt(ws.cell(row=r, column=c)).lower():
                    food_hits.append(f"{ws.title}!{ws.cell(row=r, column=c).coordinate}")
    for ws in (ov, sn):
        for row in ws.iter_rows():
            for c in row:
                if "orzo" in txt(c).lower():
                    stmt_hits.append(f"{ws.title}!{c.coordinate}")
    rec("NO ORZO in meals/ingredients", OK if not food_hits else BAD, str(food_hits))
    rec("  (orzo in rule text - expected)", "INFO",
        f"{len(stmt_hits)} rule statement(s): {stmt_hits}")

    # --- 3. Olives never on Savea's side ------------------------------------
    bad = [f"Calendar row {r}" for r in cal_rows
           if "olive" in txt(cal.cell(row=r, column=9)).lower()]
    # Snack Reference: Savea's block runs from its banner to Aaron's banner.
    s_start = a_start = None
    for r in range(1, sn.max_row + 1):
        t = txt(sn.cell(row=r, column=1)).upper()
        if "GRAB BENCH" in t:
            s_start = r
        if "AARON'S LIST" in t:
            a_start = r
    if s_start and a_start:
        for r in range(s_start + 1, a_start):
            if "olive" in txt(sn.cell(row=r, column=1)).lower():
                bad.append(f"Snack Reference row {r}")
    rec("Olives absent from Savea's bench + snack column", OK if not bad else BAD, str(bad))
    banner = txt(sn.cell(row=s_start, column=1)) if s_start else ""
    rec("  (bench labelled olive-free)", OK if "NO OLIVES" in banner.upper() else BAD, banner[:60])

    # --- 4. No dinner repeats more than twice -------------------------------
    def norm(s):
        # Drop bracketed/parenthesised qualifiers AND any trailing " - ..." clause, so
        # "Street tacos (chicken) - build-your-own bar" and "Street tacos (chicken)"
        # collapse to the same dish rather than counting as two.
        s = re.sub(r"\[.*?\]|\(.*?\)", "", s)
        s = re.split(r"\s+-\s+", s)[0]
        return re.sub(r"[^a-z ]", " ", s.lower()).split()

    dinners = []
    for r in cal_rows:
        d = txt(cal.cell(row=r, column=4))
        if not d or re.search(r"leftover|prep session", d, re.I):
            continue
        dinners.append(d)
    # Group by the first three significant words. This folds "X [DOUBLE BATCH]" and
    # "X [FROM FREEZER]" together while keeping "Greek sheet-pan chicken" distinct from
    # "Harissa sheet-pan chicken" - a similarity score merges those two, which is wrong.
    groups = collections.defaultdict(list)
    for d in dinners:
        groups[" ".join(norm(d)[:3])].append(d)
    over = {k: len(v) for k, v in groups.items() if len(v) > 2}
    rec(f"No dinner repeats >2x  ({len(dinners)} dinners, {len(groups)} unique)",
        OK if not over else BAD,
        str(over) if over else f"{sum(1 for v in groups.values() if len(v) == 2)} twice, "
                               f"{sum(1 for v in groups.values() if len(v) == 1)} once")
    for k, v in sorted(groups.items(), key=lambda x: (-len(x[1]), x[0])):
        if len(v) > 1:
            rec(f"  {len(v)}x  {k}", BAD if len(v) > 2 else "INFO", v[0])

    # --- 5. Every calendar day tagged to an order ---------------------------
    untagged = [txt(cal.cell(row=r, column=1)) for r in cal_rows
                if txt(cal.cell(row=r, column=10)) not in ("Order 1", "Order 2")]
    o1 = sum(1 for r in cal_rows if txt(cal.cell(row=r, column=10)) == "Order 1")
    o2 = sum(1 for r in cal_rows if txt(cal.cell(row=r, column=10)) == "Order 2")
    rec("Every day tagged 'Supplied by'", OK if not untagged else BAD,
        f"Order 1 = {o1} days, Order 2 = {o2} days" + (f", untagged: {untagged}" if untagged else ""))

    # --- 6. Kefir = 4 gallons, 2 per order ----------------------------------
    per = collections.Counter()
    for r in sl_rows:
        if "kefir" in txt(sl.cell(row=r, column=2)).lower():
            q = sl.cell(row=r, column=7).value or 0
            per[txt(sl.cell(row=r, column=1))] += q
    total = sum(per.values())
    rec("Kefir = 4 gallons, 2 per order",
        OK if total == 4 and set(per.values()) == {2} else BAD, f"{dict(per)} total={total}")

    # --- 7. Jerky costed in, two batches ------------------------------------
    jerky = [r for r in sl_rows if "JERKY" in txt(sl.cell(row=r, column=9)).upper()]
    both = any("BATCH 1" in txt(sl.cell(row=r, column=9)).upper()
               and "BATCH 2" in txt(sl.cell(row=r, column=9)).upper() for r in jerky) \
        or len(jerky) >= 2
    rec("Jerky beef costed, both batches", OK if jerky and both else BAD,
        f"{len(jerky)} line(s): {[txt(sl.cell(row=r, column=2)) for r in jerky]}")

    # --- 8. Dried beans only, cannellini carries the boil warning ----------
    dried = [txt(sl.cell(row=r, column=2)) for r in sl_rows
             if txt(sl.cell(row=r, column=2)).lower().startswith("dried")]
    canned = [txt(sl.cell(row=r, column=2)) for r in sl_rows
              if re.search(r"canned.*bean|bean.*canned", txt(sl.cell(row=r, column=2)), re.I)]
    rec("Beans bought dried, none canned", OK if dried and not canned else BAD,
        f"{len(dried)} dried: {dried}" + (f" | CANNED: {canned}" if canned else ""))

    cann = [r for r in sl_rows if "cannellini" in txt(sl.cell(row=r, column=2)).lower()]
    warned = all(re.search(r"10.?min|rolling boil", txt(sl.cell(row=r, column=10)), re.I)
                 for r in cann) if cann else True
    rec("Cannellini carries the 10-min rolling-boil warning",
        OK if warned else BAD, "no cannellini this month" if not cann else f"{len(cann)} line(s)")

    # --- 9. Budget lands under target ---------------------------------------
    vals = openpyxl.load_workbook(path, data_only=True)["Shopping Lists"]
    tot = budget = None
    for r in range(1, vals.max_row + 1):
        lbl = txt(vals.cell(row=r, column=2)).upper()
        v = vals.cell(row=r, column=8).value
        if "COUNTED TOWARD BUDGET" in lbl:
            tot = v
        if lbl == "MONTHLY BUDGET":
            budget = v
    if tot is None or budget is None:
        rec("Budget under target", "INFO", "run recalc.py first - formulas have no cached values yet")
    else:
        rec("Budget under target", OK if tot <= budget else BAD,
            f"${tot:,.2f} of ${budget:,.2f}  (cushion ${budget - tot:,.2f})")

    # --- 10. Totals are live formulas, not constants ------------------------
    hard = [f"H{r}" for r in sl_rows
            if not str(sl.cell(row=r, column=8).value or "").startswith("=")]
    rec("Line totals are formulas, not constants", OK if not hard else BAD,
        f"{len(hard)} hardcoded" if hard else f"{len(sl_rows)} formula cells")

    # --- 11. Cross-check against the source data file -----------------------
    if data_path:
        import importlib.util
        spec = importlib.util.spec_from_file_location("pd_", data_path)
        m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
        src = sum(x[5] * x[6] for x in (m.ORDER1 + m.ORDER2) if x[4] in ("Core", "Pantry"))
        if tot is not None:
            rec("Sheet total matches source data", OK if abs(src - tot) < 0.01 else BAD,
                f"source ${src:,.2f} vs sheet ${tot:,.2f}")
    return finish()


def finish():
    w = max(len(n) for n, _, _ in results) + 2
    fails = 0
    print()
    for n, s, d in results:
        mark = {OK: "PASS", BAD: "FAIL", "INFO": "info"}[s]
        print(f"  [{mark}] {n:<{w}} {d}")
        fails += s == BAD
    print(f"\n  {len(results)} checks, {fails} failing\n")
    return 1 if fails else 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("usage: audit_workbook.py <workbook.xlsx> [plan_data.py]")
    sys.exit(audit(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None))
